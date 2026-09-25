import io

import pandas as pd
from agents import function_tool
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from agent_server.tools.list_files import VOLUME_URI
from agent_server.utils import get_user_workspace_client


@function_tool
def parse_pptx(filename: str) -> str:
    """Parse a PPTX file from the energy reports volume and return the text blocks, one per slide.

    Use a file name returned by get_files_in_volume, not a full path.
    Each block includes slide text, chart data (as formatted tables),
    and presenter notes.
    """
    if "/" in filename:
        return "Pass only the file name, without a path."

    # Runs on behalf of the end user, so the user needs READ VOLUME on the volume
    client = get_user_workspace_client()
    response = client.files.download(f"{VOLUME_URI}/{filename}")
    if response.contents is None:
        return f"The file {filename} is empty."
    prs = Presentation(io.BytesIO(response.contents.read()))
    slides = []

    for i, slide in enumerate(prs.slides, 1):
        parts = [f"--- Slide {i} ---"]

        # --- text content ---
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        parts.append(text)

        # --- chart data ---
        chart_num = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.CHART:
                chart_num += 1
                chart = shape.chart
                title = (
                    chart.chart_title.text_frame.text
                    if chart.has_title
                    else f"Chart {chart_num}"
                )
                chart_type = str(chart.chart_type)
                plot = chart.plots[0]
                categories = [str(c) for c in plot.categories]

                series_data = {}
                for si, series in enumerate(plot.series):
                    series_data[f"Series_{si}"] = list(series.values)

                n_rows = len(next(iter(series_data.values())))
                df = pd.DataFrame(series_data, index=categories[:n_rows])
                df.index.name = "Category"

                parts.append(f"\n[{title} | {chart_type}]")
                parts.append(df.to_string())

        # --- presenter notes ---
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"\n[Presenter Notes]\n{notes}")

        slides.append("\n".join(parts))

    return "\n".join(slides)
