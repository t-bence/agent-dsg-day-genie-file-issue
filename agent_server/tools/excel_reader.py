import asyncio
import io
import re
import warnings
from dataclasses import dataclass
from typing import Any

import openpyxl
from agents import function_tool
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook

from agent_server.tools.list_files import VOLUME_URI
from agent_server.utils import get_user_workspace_client

MAX_CACHED_WORKBOOKS = 2
MAX_RANGE_CELLS = 1500
MAX_FIND_RESULTS = 50
MAX_TRACE_NODES = 150
MAX_EXPANDED_RANGE_CELLS = 30

CELL_RE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)$")
RANGE_RE = re.compile(r"^\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+$")
EXTERNAL_RE = re.compile(r"\[(\d+)\]")


@dataclass
class CachedWorkbook:
    last_modified: str | None
    formulas: Workbook
    values: Workbook
    external_ref_counts: dict[int, int] | None = None


_cache: dict[str, CachedWorkbook] = {}
_load_lock = asyncio.Lock()


def _parse_workbooks(data: bytes) -> tuple[Workbook, Workbook]:
    # openpyxl warns about unsupported Excel extensions (data validation, conditional formatting)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        # data_only=True returns the values that Excel stored when the file was last saved
        values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    return formulas, values


async def _get_workbook(filename: str) -> CachedWorkbook:
    if "/" in filename:
        raise ValueError("Pass only the file name, without a path.")
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError(f"Unsupported file type: {filename}. Supported types: .xlsx, .xlsm.")

    path = f"{VOLUME_URI}/{filename}"
    # Runs on behalf of the end user on every call, so the cache never skips the user's access check
    client = get_user_workspace_client()
    metadata = await asyncio.to_thread(client.files.get_metadata, path)

    async with _load_lock:
        cached = _cache.get(filename)
        if cached is not None and cached.last_modified == metadata.last_modified:
            return cached

        response = await asyncio.to_thread(client.files.download, path)
        if response.contents is None:
            raise ValueError(f"The file {filename} is empty.")
        data = await asyncio.to_thread(response.contents.read)
        formulas, values = await asyncio.to_thread(_parse_workbooks, data)

        _cache.pop(filename, None)
        while len(_cache) >= MAX_CACHED_WORKBOOKS:
            _cache.pop(next(iter(_cache)))
        cached = CachedWorkbook(metadata.last_modified, formulas, values)
        _cache[filename] = cached
        return cached


def _formula_text(raw: Any) -> str | None:
    """Return the formula of a cell, or None for a plain input value."""
    text = getattr(raw, "text", raw)  # ArrayFormula keeps the formula in .text
    if isinstance(text, str) and text.startswith("="):
        return text
    return None


def _format_value(value: Any) -> str:
    if value is None:
        return "(empty)"
    if isinstance(value, float):
        return f"{value:.10g}"
    text = str(value)
    return text if len(text) <= 300 else text[:300] + "..."


def _quote_sheet(sheet: str) -> str:
    return f"'{sheet}'"


def _get_sheet(wb: Workbook, sheet: str):
    if sheet not in wb.sheetnames:
        raise ValueError(f"Sheet {sheet!r} not found. Sheets: {', '.join(wb.sheetnames)}")
    return wb[sheet]


def _describe_cell(cached: CachedWorkbook, sheet: str, coordinate: str) -> str:
    formula = _formula_text(cached.formulas[sheet][coordinate].value)
    value = _format_value(cached.values[sheet][coordinate].value)
    return f"{value} | {formula}" if formula else f"{value} | input value"


def _count_external_refs(cached: CachedWorkbook) -> dict[int, int]:
    """Count the cell formulas that reference each external workbook, like [3]Sheet!A1."""
    if cached.external_ref_counts is None:
        counts: dict[int, int] = {}
        for ws in cached.formulas.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    formula = _formula_text(cell.value)
                    if formula and "[" in formula:
                        for index in set(EXTERNAL_RE.findall(formula)):
                            counts[int(index)] = counts.get(int(index), 0) + 1
        cached.external_ref_counts = counts
    return cached.external_ref_counts


def _short_list(items: list[str], limit: int) -> str:
    if not items:
        return "unknown"
    extra = f" (+{len(items) - limit} more)" if len(items) > limit else ""
    return ", ".join(items[:limit]) + extra


def _external_links(wb: Workbook) -> list:
    # openpyxl keeps external links in a private attribute and has no public API for them
    return getattr(wb, "_external_links", [])


def _bounds(cell_range: str) -> tuple[int, int, int, int]:
    """Return (min_col, min_row, max_col, max_row) for a range like B5:O30."""
    bounds = range_boundaries(cell_range.replace("$", ""))
    if any(b is None for b in bounds):
        raise ValueError(f"{cell_range!r} is a whole row or column. Use a range with start and end cells, like B5:O30.")
    return bounds  # type: ignore[return-value]


def _defined_names(wb: Workbook) -> dict[str, str]:
    names = {name: dn.attr_text for name, dn in wb.defined_names.items()}
    for ws in wb.worksheets:
        for name, dn in ws.defined_names.items():
            names[f"{ws.title}!{name}"] = dn.attr_text
    return names


@function_tool
async def excel_workbook_overview(filename: str) -> str:
    """Describe an Excel workbook from the energy reports volume.

    Returns the workbook properties, every sheet (including hidden and very hidden sheets)
    with its used range and hidden rows and columns, the defined names, and the external
    workbooks it links to, with how many formulas and defined names use each link.
    Call this first for a new workbook. Use a file name returned by get_files_in_volume.
    """
    cached = await _get_workbook(filename)
    wb = cached.formulas
    props = wb.properties
    lines = [
        f"Workbook: {filename}",
        (
            f"Title: {props.title} | Subject: {props.subject} | Creator: {props.creator} | "
            f"Last modified by: {props.lastModifiedBy} | Created: {props.created} | Modified: {props.modified}"
        ),
        "",
        f"Sheets ({len(wb.worksheets)}), in workbook order:",
    ]
    for ws in wb.worksheets:
        hidden_rows = sum(1 for dim in ws.row_dimensions.values() if dim.hidden)
        hidden_cols = sum(1 for dim in ws.column_dimensions.values() if dim.hidden)
        hidden_info = f", hidden rows: {hidden_rows}, hidden column groups: {hidden_cols}" if hidden_rows or hidden_cols else ""
        lines.append(f"- {ws.title!r}: state={ws.sheet_state}, used range={ws.dimensions}{hidden_info}")

    names = _defined_names(wb)
    ref_counts = await asyncio.to_thread(_count_external_refs, cached)
    lines += ["", f"External workbook links ({len(_external_links(wb))}):"]
    if not _external_links(wb):
        lines.append("- none")
    for index, link in enumerate(_external_links(wb), 1):
        target = link.file_link.Target if link.file_link is not None else "(unknown target)"
        book = link.externalBook
        sheet_names = list(book.sheetNames.sheetName) if book.sheetNames else []
        name_uses = sum(1 for ref in names.values() if ref and f"[{index}]" in ref)
        lines.append(
            f"- [{index}] {target} | sheets: {_short_list(sheet_names, 10)} | "
            f"used by {ref_counts.get(index, 0)} cell formulas and {name_uses} defined names"
        )

    external = {n for n, ref in names.items() if ref and EXTERNAL_RE.search(ref)}
    broken = {n for n, ref in names.items() if n not in external and ref and "#REF!" in ref}
    ranges = {n: ref for n, ref in names.items() if n not in external | broken and ref and "!" in ref}
    constants = len(names) - len(external) - len(broken) - len(ranges)
    lines += [
        "",
        (
            f"Defined names: {len(names)} in total. {len(external)} point to external workbooks, "
            f"{len(broken)} are broken (#REF!), {constants} are constants or formulas, "
            f"{len(ranges)} point to cells in this workbook:"
        ),
    ]
    lines += [f"- {name} = {ref}" for name, ref in list(ranges.items())[:50]]
    if len(ranges) > 50:
        lines.append(f"(+{len(ranges) - 50} more)")
    return "\n".join(lines)


@function_tool
async def excel_find(filename: str, text: str, sheet: str | None = None) -> str:
    """Find cells whose text contains the given phrase (case-insensitive), for example a region name or a row label.

    Searches all sheets, including hidden ones, unless a sheet name is given.
    Returns up to 50 matches as 'Sheet'!Cell: text. Use this to locate labels before reading a range.
    """
    cached = await _get_workbook(filename)
    wb = cached.values
    sheets = [_get_sheet(wb, sheet)] if sheet else wb.worksheets
    needle = text.lower()
    matches = []
    for ws in sheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and needle in cell.value.lower():
                    matches.append(f"{_quote_sheet(ws.title)}!{cell.coordinate}: {_format_value(cell.value)}")
                    if len(matches) >= MAX_FIND_RESULTS:
                        return "\n".join(matches) + f"\n(Stopped at {MAX_FIND_RESULTS} matches. Narrow the search.)"
    return "\n".join(matches) if matches else f"No cells contain {text!r}."


@function_tool
async def excel_read_range(filename: str, sheet: str, cell_range: str) -> str:
    """Read the cells in a range of one sheet, for example 'B5:O30'. Hidden sheets work too.

    For every non-empty cell, returns the stored value and, for formula cells, the formula.
    Output lines look like: H11: 191.97 | ='1b Direct Fuel Cost Component'!H62
    Rows that are hidden in Excel are marked. Reads at most 1500 cells per call.
    """
    cached = await _get_workbook(filename)
    ws_formulas = _get_sheet(cached.formulas, sheet)
    ws_values = cached.values[sheet]
    min_col, min_row, max_col, max_row = _bounds(cell_range)

    lines = [f"{_quote_sheet(sheet)}!{cell_range}"]
    count = 0
    for row_index in range(min_row, max_row + 1):
        hidden = ws_formulas.row_dimensions[row_index].hidden
        for col_index in range(min_col, max_col + 1):
            formula = _formula_text(ws_formulas.cell(row_index, col_index).value)
            value = ws_values.cell(row_index, col_index).value
            if value is None and formula is None:
                continue
            coordinate = f"{get_column_letter(col_index)}{row_index}"
            line = f"{coordinate}: {_format_value(value)}"
            if formula:
                line += f" | {formula}"
            if hidden:
                line += " | (hidden row)"
            lines.append(line)
            count += 1
            if count >= MAX_RANGE_CELLS:
                lines.append(f"(Stopped at {MAX_RANGE_CELLS} cells, after row {row_index}. Read a smaller range.)")
                return "\n".join(lines)
    if count == 0:
        lines.append("(all cells are empty)")
    return "\n".join(lines)


def _split_reference(reference: str, current_sheet: str) -> tuple[str | None, str, str]:
    """Split a reference like 'Sheet 1'!$A$1 into (external index, sheet, address)."""
    if "!" in reference:
        sheet_part, address = reference.rsplit("!", 1)
        sheet_part = sheet_part.strip("'").replace("''", "'")
    else:
        sheet_part, address = current_sheet, reference
    external = None
    match = re.match(r"^\[(\d+)\](.*)$", sheet_part)
    if match:
        external, sheet_part = match.group(1), match.group(2)
    return external, sheet_part, address.replace("$", "")


def _formula_references(formula: str) -> list[str]:
    tokens = Tokenizer(formula).items
    references = [t.value for t in tokens if t.type == Token.OPERAND and t.subtype == Token.RANGE]
    # A formula like IF(H16="-","-",H16*2) names the same cell twice; trace it once
    return list(dict.fromkeys(references))


def _trace(cached: CachedWorkbook, sheet: str, coordinate: str, max_depth: int) -> str:
    names = _defined_names(cached.formulas)
    lines: list[str] = []
    seen: set[tuple[str, str]] = set()

    def visit_reference(reference: str, current_sheet: str, depth: int) -> None:
        if len(lines) >= MAX_TRACE_NODES:
            return
        indent = "  " * depth
        external, ref_sheet, address = _split_reference(reference, current_sheet)
        if external is not None:
            link = _external_links(cached.formulas)[int(external) - 1]
            target = link.file_link.Target if link.file_link is not None else "unknown file"
            lines.append(f"{indent}{reference} -> external workbook [{external}] {target} (cannot be traced further)")
            return
        if address in names or f"{ref_sheet}!{address}" in names:
            target = names.get(f"{ref_sheet}!{address}") or names[address]
            lines.append(f"{indent}defined name {address} = {target}")
            for inner in _formula_references("=" + target):
                visit_reference(inner, ref_sheet, depth + 1)
            return
        if ref_sheet not in cached.formulas.sheetnames:
            lines.append(f"{indent}{reference} -> unknown sheet or name")
            return
        if CELL_RE.match(address):
            visit_cell(ref_sheet, address, depth)
            return
        if RANGE_RE.match(address):
            min_col, min_row, max_col, max_row = _bounds(address)
            size = (max_col - min_col + 1) * (max_row - min_row + 1)
            label = f"{_quote_sheet(ref_sheet)}!{address}"
            if size > MAX_EXPANDED_RANGE_CELLS:
                lines.append(f"{indent}{label} -> range of {size} cells, not expanded (read it or trace single cells)")
                return
            lines.append(f"{indent}{label} -> range of {size} cells:")
            for row_index in range(min_row, max_row + 1):
                for col_index in range(min_col, max_col + 1):
                    visit_cell(ref_sheet, f"{get_column_letter(col_index)}{row_index}", depth + 1)
            return
        lines.append(f"{indent}{reference} -> whole row or column reference, not expanded")

    def visit_cell(cell_sheet: str, cell_coordinate: str, depth: int) -> None:
        if len(lines) >= MAX_TRACE_NODES:
            return
        indent = "  " * depth
        label = f"{_quote_sheet(cell_sheet)}!{cell_coordinate}"
        if (cell_sheet, cell_coordinate) in seen:
            lines.append(f"{indent}{label} (already traced above)")
            return
        seen.add((cell_sheet, cell_coordinate))
        lines.append(f"{indent}{label} = {_describe_cell(cached, cell_sheet, cell_coordinate)}")
        formula = _formula_text(cached.formulas[cell_sheet][cell_coordinate].value)
        if formula is None:
            return
        if depth >= max_depth:
            lines.append(f"{indent}  (max depth {max_depth} reached, trace this cell again to go deeper)")
            return
        for reference in _formula_references(formula):
            visit_reference(reference, cell_sheet, depth + 1)

    visit_cell(sheet, coordinate, 0)
    if len(lines) >= MAX_TRACE_NODES:
        lines.append(f"(Stopped at {MAX_TRACE_NODES} lines. Trace a deeper cell again to continue.)")
    return "\n".join(lines)


@function_tool
async def excel_trace_precedents(filename: str, sheet: str, cell: str, max_depth: int = 10) -> str:
    """Trace a formula cell back to its inputs, across sheets, like Excel's Trace Precedents.

    Returns an indented tree. Each line shows the cell, its stored value and its formula,
    or 'input value' when the cell is a plain input. Defined names are resolved,
    small ranges are expanded, and references to external workbooks are marked.
    """
    cached = await _get_workbook(filename)
    _get_sheet(cached.formulas, sheet)
    coordinate = cell.replace("$", "").upper()
    if not CELL_RE.match(coordinate):
        return f"{cell!r} is not a single cell address like H11."
    return await asyncio.to_thread(_trace, cached, sheet, coordinate, max_depth)
